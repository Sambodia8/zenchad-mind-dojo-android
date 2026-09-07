from __future__ import annotations

import json
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "mark-yoga-voice-recording-script.docx"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
HEADING_BLUE = RGBColor(0x2E, 0x74, 0xB5)
HEADING_DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK_BLUE = RGBColor(0x0B, 0x25, 0x45)
MUTED = RGBColor(0x5F, 0x69, 0x74)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"


CLASS_INTROS = {
    "sun-salutation": (
        "Welcome to Sun Salutation with Mark. Move with your breath, stay within a comfortable "
        "range, and use a knees-down option whenever you need it."
    ),
    "daily-reset": (
        "Welcome to Daily Reset with Mark. We will move calmly through the whole body. Keep every "
        "stretch comfortable and let your breathing stay easy."
    ),
    "before-run": (
        "Welcome to Before Running with Mark. Start gently and let each movement warm your legs "
        "without forcing the range."
    ),
    "before-cycling": (
        "Welcome to Before Cycling with Mark. We will prepare the hips, thighs and knees before an "
        "easy, progressive warm-up on the bike."
    ),
    "after-run": (
        "Welcome to After Running with Mark. Let your breathing settle, then move through these "
        "stretches gently and without bouncing."
    ),
    "after-cycling": (
        "Welcome to After Cycling with Mark. Ease down from standing to the floor and keep every "
        "stretch steady, relaxed and pain-free."
    ),
    "back-and-shoulders": (
        "Welcome to Back and Shoulders with Mark. Move slowly through a comfortable range, and stop "
        "if pain becomes sharp, travels or causes tingling."
    ),
    "gentle-leg-recovery": (
        "Welcome to Gentle Leg Recovery with Mark. This class is only for mild, recovering soreness. "
        "Stop if pain worsens, and do not use it for a new or acute injury."
    ),
}


SYSTEM_LINES = [
    ("yoga-system-transition", "yoga-system-transition.wav", "Get ready for the next pose."),
    ("yoga-system-switch-sides", "yoga-system-switch-sides.wav", "Switch sides."),
    ("yoga-system-pause", "yoga-system-pause.wav", "Pause here whenever you need to."),
    ("yoga-system-resume", "yoga-system-resume.wav", "When you are ready, let's continue."),
    (
        "yoga-system-complete",
        "yoga-system-complete.wav",
        "Class complete. Take one easy breath and notice how you feel.",
    ),
]


def load_yoga_data() -> dict:
    node = shutil.which("node")
    if not node:
        raise RuntimeError("Node.js is required to read the canonical Yoga data.")
    code = """
import { YOGA_CLASSES, expandYogaClassSlides } from './src/data.ts';
const targetClassIds = new Set([
  'sun-salutation',
  'daily-reset',
  'before-run',
  'before-cycling',
  'after-run',
  'after-cycling',
  'back-and-shoulders',
  'gentle-leg-recovery'
]);
const targetClasses = YOGA_CLASSES.filter((yogaClass) => targetClassIds.has(yogaClass.id));
const movements = new Map();
for (const yogaClass of targetClasses) {
  for (const slide of expandYogaClassSlides(yogaClass)) {
    if (!movements.has(slide.movement.id)) {
      movements.set(slide.movement.id, {
        id: slide.movement.id,
        name: slide.movement.name,
        cue: slide.movement.cue,
        sensationCue: slide.movement.sensationCue,
        muscleGroups: slide.movement.muscleGroups,
        sensationKind: slide.movement.sensationKind
      });
    }
  }
}
console.log(JSON.stringify({
  classes: targetClasses.map(({ id, name, timing, description, safetyGate }) => ({
    id, name, timing, description, safetyGate: Boolean(safetyGate)
  })),
  movements: [...movements.values()]
}));
"""
    result = subprocess.run(
        [
            node,
            "--no-warnings",
            "--experimental-strip-types",
            "--experimental-loader",
            "./scripts/ts-loader.mjs",
            "--input-type=module",
            "-e",
            code,
        ],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    return json.loads(result.stdout)


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, HEADING_BLUE, 18, 10),
        "Heading 2": (13, HEADING_BLUE, 14, 7),
        "Heading 3": (12, HEADING_DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    clip_style = document.styles.add_style("Clip Label", 1)
    clip_style.font.name = "Calibri"
    clip_style.font.size = Pt(8.5)
    clip_style.font.bold = True
    clip_style.font.color.rgb = MUTED
    clip_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    clip_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    clip_style.paragraph_format.space_before = Pt(0)
    clip_style.paragraph_format.space_after = Pt(2)
    clip_style.paragraph_format.keep_with_next = True

    script_style = document.styles.add_style("Script Line", 1)
    script_style.font.name = "Calibri"
    script_style.font.size = Pt(11)
    script_style.font.color.rgb = INK_BLUE
    script_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    script_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    script_style.paragraph_format.left_indent = Inches(0.18)
    script_style.paragraph_format.right_indent = Inches(0.1)
    script_style.paragraph_format.space_before = Pt(0)
    script_style.paragraph_format.space_after = Pt(7)
    script_style.paragraph_format.line_spacing = 1.25


def set_cell_shading(cell, fill: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        cell_properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")

    indent = table_properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        table_properties.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.first_child_found_in("w:tcW")
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.append(cell_width)
            cell_width.set(qn("w:w"), str(widths[index]))
            cell_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)
    set_run_font(run, size=9, color=MUTED)


def add_running_furniture(document: Document) -> None:
    section = document.sections[0]
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("ZENCHAD YOGA  |  MARK RECORDING SCRIPT")
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    label = paragraph.add_run("Page ")
    set_run_font(label, size=9, color=MUTED)
    add_page_field(paragraph)


def add_title_block(document: Document, class_count: int, movement_count: int) -> None:
    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(8)
    kicker.paragraph_format.space_after = Pt(2)
    run = kicker.add_run("ZENCHAD YOGA")
    set_run_font(run, size=10, color=HEADING_BLUE, bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(7)
    run = title.add_run("Mark Voice Recording Script")
    set_run_font(run, size=28, color=INK_BLUE, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run(
        f"Studio master | {class_count} classes | {movement_count} reusable movements | Version 1"
    )
    set_run_font(run, size=11.5, color=MUTED)

    lead = document.add_paragraph()
    lead.paragraph_format.space_after = Pt(12)
    run = lead.add_run(
        "Record each filename as a separate clean take. Pose wording mirrors the canonical app "
        "guidance so the finished audio can be mapped without rewriting or duplication."
    )
    set_run_font(run, size=11, color=INK_BLUE, bold=True)


def add_spec_table(document: Document) -> None:
    document.add_heading("Recording setup", level=1)
    rows = [
        ("Format", "48 kHz / 24-bit mono WAV, clean and dry"),
        ("Files", "One file per listed line; use the filename exactly as written"),
        ("Room tone", "Leave roughly 0.5 seconds of clean room tone before and after each take"),
        ("Processing", "No music, reverb, noise gate, compression or loudness normalisation"),
        ("Delivery", "Warm, calm and direct; sound like a coach beside the listener, not an announcer"),
        ("Pace", "Unhurried but concise; allow natural breath-sized pauses at full stops"),
        ("Pickups", "Record a complete replacement take when a line needs correcting"),
    ]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1700, CONTENT_WIDTH_DXA - 1700])
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Studio requirement"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, LIGHT_BLUE)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9.5, color=INK_BLUE, bold=True)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        for run in cells[0].paragraphs[0].runs:
            set_run_font(run, size=9.5, color=INK_BLUE, bold=True)
        for run in cells[1].paragraphs[0].runs:
            set_run_font(run, size=9.5, color=INK_BLUE)
        set_table_geometry(table, [1700, CONTENT_WIDTH_DXA - 1700])


def add_performance_direction(document: Document) -> None:
    document.add_heading("Performance direction", level=1)
    directions = [
        ("Tone", "Grounded, reassuring and human. Avoid a theatrical meditation voice."),
        ("Energy", "Clear enough to guide movement while remaining calm over a soft soundtrack."),
        ("Instruction", "Use gentle emphasis on the action words; never sound urgent or corrective."),
        ("Safety", "Deliver warnings plainly and without alarm. Do not soften or improvise their meaning."),
        ("Pronunciation", "Say ZenChad as 'Zen-chad'. Read NSDR letter by letter if it appears in later pickups."),
    ]
    for label, copy in directions:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        lead = paragraph.add_run(f"{label}: ")
        set_run_font(lead, color=INK_BLUE, bold=True)
        body = paragraph.add_run(copy)
        set_run_font(body, color=INK_BLUE)


def add_script_entry(
    document: Document,
    label: str,
    filename: str,
    script: str,
    *,
    keep_with_next: bool = False,
) -> None:
    meta = document.add_paragraph(style="Clip Label")
    meta.add_run(f"{label.upper()}  |  {filename}")
    line = document.add_paragraph(style="Script Line")
    line.paragraph_format.keep_with_next = keep_with_next
    line.add_run(script)


def add_class_intros(document: Document, classes: list[dict]) -> None:
    document.add_heading("Class introductions", level=1)
    intro = document.add_paragraph(
        "Record once per built-in class. These lines establish the purpose and safety tone before the first pose."
    )
    intro.paragraph_format.space_after = Pt(9)
    for yoga_class in classes:
        script = CLASS_INTROS[yoga_class["id"]]
        heading = document.add_heading(yoga_class["name"], level=2)
        heading.paragraph_format.keep_with_next = True
        add_script_entry(
            document,
            "Class intro",
            f"yoga-class-{yoga_class['id']}-intro.wav",
            script,
        )


def add_system_lines(document: Document) -> None:
    document.add_heading("Shared player lines", level=1)
    intro = document.add_paragraph(
        "These clips are reusable across every built-in or custom class. Keep them short and neutral."
    )
    intro.paragraph_format.space_after = Pt(9)
    for clip_id, filename, script in SYSTEM_LINES:
        heading = document.add_heading(clip_id.replace("yoga-system-", "").replace("-", " ").title(), level=2)
        heading.paragraph_format.keep_with_next = True
        add_script_entry(document, "Shared line", filename, script)


def add_pose_lines(document: Document, movements: list[dict]) -> None:
    document.add_heading("Reusable pose lines", level=1)
    intro = document.add_paragraph(
        "The start line names the pose and repeats the app's setup cue. The optional detail line mirrors "
        "the hidden Pose guidance sheet and can be played only when extra guidance is requested."
    )
    intro.paragraph_format.space_after = Pt(9)
    for movement in movements:
        heading = document.add_heading(movement["name"], level=2)
        heading.paragraph_format.keep_with_next = True
        start_script = f"{movement['name']}. {movement['cue']}"
        add_script_entry(
            document,
            "Pose start",
            f"yoga-pose-{movement['id']}-start.wav",
            start_script,
            keep_with_next=True,
        )
        detail_label = "Optional stretch detail" if movement["sensationKind"] == "stretch" else "Optional working detail"
        add_script_entry(
            document,
            detail_label,
            f"yoga-pose-{movement['id']}-detail.wav",
            movement["sensationCue"],
            keep_with_next=True,
        )
        muscles = document.add_paragraph(style="Clip Label")
        muscles.paragraph_format.space_after = Pt(8)
        muscles.paragraph_format.keep_with_next = False
        muscles.add_run(f"REFERENCE MUSCLES  |  {', '.join(movement['muscleGroups'])}")


def build_document() -> Path:
    data = load_yoga_data()
    if len(data["classes"]) != 8 or len(data["movements"]) != 28:
        raise RuntimeError(
            f"Expected 8 classes and 28 movements, found {len(data['classes'])} and {len(data['movements'])}."
        )

    document = Document()
    document.settings.odd_and_even_pages_header_footer = False
    section = document.sections[0]
    section.different_first_page_header_footer = False
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_styles(document)
    add_running_furniture(document)
    add_title_block(document, len(data["classes"]), len(data["movements"]))
    add_spec_table(document)
    add_performance_direction(document)
    add_class_intros(document, data["classes"])
    add_system_lines(document)
    add_pose_lines(document, data["movements"])

    core = document.core_properties
    core.title = "ZenChad Yoga - Mark Voice Recording Script"
    core.subject = "Studio recording pack for ZenChad Yoga guidance"
    core.author = "ZenChad"
    core.keywords = "ZenChad, Yoga, Mark, voice recording, studio script"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
